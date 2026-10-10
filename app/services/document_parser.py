from __future__ import annotations

import io
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".txt", ".md", ".html", ".htm", ".pdf", ".docx", ".pptx", ".xlsx"
}


def extract_text_from_file(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))
        raise ValueError(f"不支持的文件类型：{suffix or 'unknown'}，当前支持 {supported}")
    if suffix in {".txt", ".md"}:
        return _decode_text(data)
    if suffix in {".html", ".htm"}:
        return _extract_html_text(data)
    if suffix == ".pdf":
        text, _ = extract_pdf_document(data)
        return text
    if suffix == ".docx":
        return _extract_docx_text(data)
    if suffix == ".pptx":
        return _extract_pptx_text(data)
    if suffix == ".xlsx":
        return _extract_xlsx_text(data)
    return ""


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.blocks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip_depth += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.blocks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.blocks.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth and data.strip():
            self.blocks.append(data.strip() + " ")


def _extract_html_text(data: bytes) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(_decode_text(data))
    return "\n".join(line.strip() for line in "".join(parser.blocks).splitlines() if line.strip())


def _extract_pdf_text_pypdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"\n\n[Page {index}]\n{text.strip()}")
    return "\n".join(pages).strip()


def _markdown_table(rows: list[list[Any]]) -> str:
    normalized = [
        [str(cell or "").replace("|", "\\|").replace("\n", " ").strip() for cell in row]
        for row in rows
        if row
    ]
    if not normalized:
        return ""
    width = max(len(row) for row in normalized)
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    header = normalized[0]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _body_text_without_tables(page: Any, table_boxes: list[tuple[float, ...]]) -> str:
    """Read text blocks in page order, excluding blocks already captured as tables."""

    body: list[str] = []
    for block in page.get_text("blocks", sort=True):
        if len(block) > 6 and block[6] != 0:
            continue
        x0, y0, x1, y1 = block[:4]
        center_x, center_y = (x0 + x1) / 2, (y0 + y1) / 2
        if any(
            left <= center_x <= right and top <= center_y <= bottom
            for left, top, right, bottom in table_boxes
        ):
            continue
        value = str(block[4]).strip()
        if value:
            body.append(value)
    return "\n\n".join(body)


def _needs_ocr(body_text: str, table_texts: list[str], image_count: int) -> bool:
    visible_length = len(body_text) + sum(len(table) for table in table_texts)
    return visible_length < 40 or (image_count > 0 and visible_length < 160)


def _table_context(page: Any, table_box: tuple[float, ...]) -> str:
    """Keep the heading and nearby rule that introduce a detected table."""

    preceding = [
        str(block[4]).strip()
        for block in page.get_text("blocks", sort=True)
        if len(block) > 6 and block[6] == 0
        and block[3] <= table_box[1]
        and str(block[4]).strip()
    ]
    if not preceding:
        return ""
    heading = re.compile(r"^(?:第[一二三四五六七八九十百零〇0-9]+条|附件\s*$)")
    start = next(
        (index for index in range(len(preceding) - 1, -1, -1)
         if heading.match(preceding[index])),
        max(0, len(preceding) - 3),
    )
    return " ".join(preceding[start:])[-300:]


def extract_pdf_document(data: bytes) -> tuple[str, dict[str, Any]]:
    """Extract PDF tables/images before body text and OCR low-text pages on demand."""

    try:
        import fitz
    except ImportError:
        return _extract_pdf_text_pypdf(data), {
            "loader": "pypdf",
            "layout_analysis": False,
            "ocr_available": False,
        }

    document = fitz.open(stream=data, filetype="pdf")
    blocks: list[str] = []
    page_metadata: list[dict[str, Any]] = []
    for page_index, page in enumerate(document, start=1):
        page_blocks: list[str] = []
        table_texts: list[str] = []
        table_count = 0
        table_boxes: list[tuple[float, ...]] = []
        try:
            finder = page.find_tables()
            for table_index, table in enumerate(getattr(finder, "tables", []), start=1):
                markdown = _markdown_table(table.extract())
                if markdown:
                    context = _table_context(page, tuple(table.bbox))
                    if context:
                        heading, separator, *rows = markdown.splitlines()
                        markdown = "\n".join(
                            [heading, separator, f"[Context] {context}", *rows]
                        )
                    page_blocks.append(f"[Table {page_index}-{table_index}]\n{markdown}")
                    table_texts.append(markdown)
                    table_boxes.append(tuple(table.bbox))
                    table_count += 1
        except Exception:
            table_count = 0

        body_text = _body_text_without_tables(page, table_boxes)
        image_count = len(page.get_images(full=True))
        needs_ocr = _needs_ocr(body_text, table_texts, image_count)
        ocr_applied = False
        ocr_text = ""
        ocr_confidence: float | None = None
        if needs_ocr:
            try:
                import pytesseract
                from PIL import Image

                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                ocr_text = pytesseract.image_to_string(image, lang="chi_sim+eng").strip()
                ocr_data = pytesseract.image_to_data(
                    image,
                    lang="chi_sim+eng",
                    output_type=pytesseract.Output.DICT,
                )
                confidences = [
                    float(value)
                    for value in ocr_data.get("conf", [])
                    if str(value).replace(".", "", 1).lstrip("-").isdigit() and float(value) >= 0
                ]
                if confidences:
                    ocr_confidence = sum(confidences) / len(confidences) / 100
                ocr_applied = bool(ocr_text)
            except Exception:
                ocr_applied = False
        if ocr_text and ocr_text not in body_text:
            page_blocks.append(f"[OCR page={page_index}]\n{ocr_text}")
        if body_text:
            page_blocks.append(body_text)
        if page_blocks:
            blocks.append(f"[Page {page_index}]\n" + "\n\n".join(page_blocks))
        page_metadata.append(
            {
                "page_number": page_index,
                "table_count": table_count,
                "image_count": image_count,
                "ocr_required": needs_ocr,
                "ocr_applied": ocr_applied,
                "ocr_confidence": ocr_confidence,
                "needs_review": (
                    needs_ocr
                    and (
                        not ocr_applied
                        or (ocr_confidence is not None and ocr_confidence < 0.6)
                    )
                ),
            }
        )
    document.close()
    return "\n\n".join(blocks).strip(), {
        "loader": "pymupdf-layout",
        "layout_analysis": True,
        "page_count": len(page_metadata),
        "pages": page_metadata,
        "ocr_available": any(item["ocr_applied"] for item in page_metadata),
    }


def _extract_docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks).strip()


def _extract_pptx_text(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                blocks.append(shape.text.strip())
            if shape.has_table:
                table = _markdown_table(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
                if table:
                    blocks.append(table)
        if blocks:
            slides.append(f"[Slide {index}]\n" + "\n\n".join(blocks))
    return "\n\n".join(slides)


def _extract_xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for sheet in workbook.worksheets:
            rows = [
                [str(value) if value is not None else "" for value in row]
                for row in sheet.iter_rows(values_only=True)
                if any(value is not None and str(value).strip() for value in row)
            ]
            if rows:
                sheets.append(f"[Sheet {sheet.title}]\n" + _markdown_table(rows))
        return "\n\n".join(sheets)
    finally:
        workbook.close()
