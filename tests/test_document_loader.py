from __future__ import annotations

import io

import pytest

from app.services import document_loader
from app.services.document_loader import DocumentLoadError, LoadedDocument, load_document_from_file


def test_load_document_falls_back_to_simple_text_parser(monkeypatch: pytest.MonkeyPatch) -> None:
    def missing_unstructured(filename: str, data: bytes) -> LoadedDocument:
        raise ImportError("unstructured is not installed")

    monkeypatch.setattr(document_loader, "_load_with_unstructured", missing_unstructured)

    loaded = load_document_from_file("policy.txt", "员工应提前预订差旅行程。".encode("utf-8"))

    assert loaded.loader == "simple"
    assert loaded.metadata["fallback"] is True
    assert "提前预订" in loaded.text


def test_load_document_prefers_unstructured_when_available(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_unstructured(filename: str, data: bytes) -> LoadedDocument:
        return LoadedDocument(text="结构化解析结果", loader="unstructured", element_count=3)

    monkeypatch.setattr(document_loader, "_load_with_unstructured", fake_unstructured)

    loaded = load_document_from_file("policy.docx", b"docx")

    assert loaded.loader == "unstructured"
    assert loaded.element_count == 3
    assert loaded.text == "结构化解析结果"


def test_load_document_rejects_unsupported_extension() -> None:
    with pytest.raises(DocumentLoadError):
        load_document_from_file("policy.exe", b"bad")


def test_load_html_without_unstructured(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(document_loader, "_load_with_unstructured", lambda *_: (_ for _ in ()).throw(ImportError()))

    loaded = load_document_from_file(
        "policy.html", b"<h1>Travel policy</h1><script>hidden()</script><p>Book early.</p>"
    )

    assert "Travel policy" in loaded.text
    assert "Book early" in loaded.text
    assert "hidden" not in loaded.text


def test_load_pptx_without_unstructured(monkeypatch: pytest.MonkeyPatch) -> None:
    from pptx import Presentation

    monkeypatch.setattr(document_loader, "_load_with_unstructured", lambda *_: (_ for _ in ()).throw(ImportError()))
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Travel rules"
    data = io.BytesIO()
    presentation.save(data)

    loaded = load_document_from_file("policy.pptx", data.getvalue())

    assert "[Slide 1]" in loaded.text
    assert "Travel rules" in loaded.text


def test_load_xlsx_without_unstructured(monkeypatch: pytest.MonkeyPatch) -> None:
    from openpyxl import Workbook

    monkeypatch.setattr(document_loader, "_load_with_unstructured", lambda *_: (_ for _ in ()).throw(ImportError()))
    workbook = Workbook()
    workbook.active.append(["Grade", "Hotel limit"])
    workbook.active.append(["staff", 800])
    data = io.BytesIO()
    workbook.save(data)

    loaded = load_document_from_file("policy.xlsx", data.getvalue())

    assert "[Sheet Sheet]" in loaded.text
    assert "staff" in loaded.text
    assert "800" in loaded.text
